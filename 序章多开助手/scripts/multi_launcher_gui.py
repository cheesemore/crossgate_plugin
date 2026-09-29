#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""新序章多开器 — 管理账号库，一键启动并自动登录/拉取多控/一键召唤。

启动游戏后自动连接精简桥接（SeqChapterMiniBridge），依次走
workflow_step1（登录→进游戏→拉起离线多控→一键召唤）。全程靠协议
（IPC / team_num / multi_ready）判定，不依赖坐标；进度由游戏内 Tip 飘字反馈。
需先在「序章补丁」勾选「注入精简桥接」。
"""
from __future__ import annotations

import argparse
import subprocess
import sys
import threading
import time
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, simpledialog, ttk

SHARED = Path(__file__).resolve().parents[2] / "序章助手共享"
sys.path.insert(0, str(SHARED))

from assistant_common import ipc  # noqa: E402
from assistant_common.accounts import (  # noqa: E402
    AccountProfile,
    delete_account,
    load_accounts,
    normalize_secondary_code,
    upsert_account,
)
from assistant_common.config import get_game_root, load_settings, save_settings, set_game_root  # noqa: E402
from assistant_common.game import GameInstance, launch_game  # noqa: E402
from assistant_common.patch_bridge import is_mini_bridge_ready  # noqa: E402
from assistant_common.single_instance import ensure_single_instance  # noqa: E402

APP_TITLE = "新序章多开器"

# 自动分窗：每列最多 10 个，阶梯 80px；第 11 个起右移 1200px。一次最多 20 个。
TILE_STEP_PX = 80
TILE_COL_OFFSET_PX = 1200
TILE_PER_COL = 10
TILE_MAX_WINDOWS = 20
TILE_WAIT_HWND_SEC = 45


def _tile_xy(index: int) -> tuple[int, int]:
    """index 从 0 起：第 1 个 (0,0)，第 2 个 (80,80)…；第 11 个 (1200,0)。"""
    col = index // TILE_PER_COL
    row = index % TILE_PER_COL
    return col * TILE_COL_OFFSET_PX + row * TILE_STEP_PX, row * TILE_STEP_PX


def _find_hwnd_for_pid(pid: int) -> int | None:
    if sys.platform != "win32" or pid <= 0:
        return None
    import ctypes
    from ctypes import wintypes

    user32 = ctypes.windll.user32
    best_hwnd = 0
    best_score = -1

    @ctypes.WINFUNCTYPE(wintypes.BOOL, wintypes.HWND, wintypes.LPARAM)
    def callback(hwnd, _lparam):
        nonlocal best_hwnd, best_score
        proc_id = wintypes.DWORD()
        user32.GetWindowThreadProcessId(hwnd, ctypes.byref(proc_id))
        if int(proc_id.value) != pid:
            return True
        visible = bool(user32.IsWindowVisible(hwnd))
        length = user32.GetWindowTextLengthW(hwnd)
        title = ""
        if length > 0:
            buf = ctypes.create_unicode_buffer(length + 1)
            user32.GetWindowTextW(hwnd, buf, length + 1)
            title = (buf.value or "").strip()
        cls_buf = ctypes.create_unicode_buffer(256)
        user32.GetClassNameW(hwnd, cls_buf, 256)
        cls = (cls_buf.value or "").strip()
        score = 0
        if visible:
            score += 100
        if title:
            score += min(len(title), 80)
        if "Unity" in cls:
            score += 20
        if score > best_score:
            best_score = score
            best_hwnd = int(hwnd)
        return True

    user32.EnumWindows(callback, 0)
    return best_hwnd or None


def _wait_hwnd_for_pid(pid: int, timeout: float = TILE_WAIT_HWND_SEC) -> int | None:
    deadline = time.time() + timeout
    while time.time() < deadline:
        hwnd = _find_hwnd_for_pid(pid)
        if hwnd:
            return hwnd
        time.sleep(0.4)
    return None


def _move_window(hwnd: int, x: int, y: int) -> bool:
    if sys.platform != "win32" or not hwnd:
        return False
    import ctypes
    from ctypes import wintypes

    user32 = ctypes.WinDLL("user32", use_last_error=True)
    user32.ShowWindow.argtypes = [wintypes.HWND, ctypes.c_int]
    user32.ShowWindow.restype = wintypes.BOOL
    user32.SetWindowPos.argtypes = [
        wintypes.HWND,
        wintypes.HWND,
        ctypes.c_int,
        ctypes.c_int,
        ctypes.c_int,
        ctypes.c_int,
        wintypes.UINT,
    ]
    user32.SetWindowPos.restype = wintypes.BOOL
    sw_restore = 9
    swp_nosize = 0x0001
    swp_nozorder = 0x0004
    user32.ShowWindow(hwnd, sw_restore)
    return bool(user32.SetWindowPos(hwnd, 0, x, y, 0, 0, swp_nosize | swp_nozorder))


class MultiLauncherApp:
    def __init__(self) -> None:
        self.root = tk.Tk()
        self.root.title(APP_TITLE)
        self.root.geometry("900x820")
        self.root.minsize(780, 680)

        self.instances: list[GameInstance] = []
        initial_root = get_game_root()
        # 若 settings 尚未持久化有效目录，把探测到的目录存下来，
        # 与傻瓜补丁 GUI 的默认目录保持一致，避免检测错目录。
        raw = load_settings().get("game_root", "").strip()
        if not raw or not (Path(raw) / "cg37_Data").is_dir():
            try:
                set_game_root(initial_root)
            except OSError:
                pass
        self.game_root_var = tk.StringVar(value=str(initial_root))

        outer = ttk.Frame(self.root, padding=14)
        outer.pack(fill=tk.BOTH, expand=True)

        ttk.Label(outer, text=APP_TITLE, font=("Microsoft YaHei UI", 14, "bold")).pack(anchor=tk.W)
        ttk.Label(
            outer,
            text="管理账号库，一键启动游戏后自动连接精简桥接并完成 登录→拉多控→一键召唤。批量功能为协议驱动，需先注入精简桥接（序章补丁中勾选「注入精简桥接」）。",
            foreground="#555",
            wraplength=860,
        ).pack(anchor=tk.W, pady=(0, 10))

        path_frm = ttk.LabelFrame(outer, text="游戏目录", padding=8)
        path_frm.pack(fill=tk.X, pady=(0, 10))
        row = ttk.Frame(path_frm)
        row.pack(fill=tk.X)
        ttk.Entry(row, textvariable=self.game_root_var).pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Button(row, text="选择目录", command=self.pick_game_dir, width=10).pack(side=tk.LEFT, padx=(6, 0))

        batch_frm = ttk.LabelFrame(outer, text="批量一键（协议驱动，需先注入精简桥接）", padding=8)
        batch_frm.pack(fill=tk.X, pady=(0, 10))
        batch_row = ttk.Frame(batch_frm)
        batch_row.pack(fill=tk.X)
        self.batch_login_btn = ttk.Button(
            batch_row,
            text="一键登录并拉取多控",
            command=self.batch_login_fetch,
            width=22,
        )
        self.batch_login_btn.pack(side=tk.LEFT, padx=(0, 8))
        self.batch_summon_btn = ttk.Button(
            batch_row,
            text="一键召唤",
            command=self.batch_summon,
            width=14,
        )
        self.batch_summon_btn.pack(side=tk.LEFT)
        ttk.Button(
            batch_row,
            text="开启窗口监视",
            command=self.launch_window_monitor,
            width=14,
        ).pack(side=tk.LEFT, padx=(8, 0))
        self.batch_status_var = tk.StringVar(value="批量状态：就绪")
        ttk.Label(
            batch_frm,
            textvariable=self.batch_status_var,
            font=("Microsoft YaHei UI", 8),
            foreground="#888",
        ).pack(anchor=tk.W, pady=(6, 0))
        ttk.Label(
            batch_frm,
            text="「一键启动」会自动连接精简桥接并完成登录→进游戏→拉起离线多控→一键召唤（含 Tip 飘字）；"
            "下方批量按钮可按需单独触发登录/召唤，按 team≥5 协议判定聚齐（不依赖坐标）。",
            font=("Microsoft YaHei UI", 8),
            foreground="#888",
            wraplength=860,
        ).pack(anchor=tk.W, pady=(2, 0))

        acc_frm = ttk.LabelFrame(outer, text="账号库", padding=8)
        acc_frm.pack(fill=tk.BOTH, expand=True)

        cols = ("label", "phone")
        self.acc_tree = ttk.Treeview(acc_frm, columns=cols, show="headings", selectmode="extended")
        self.acc_tree.heading("label", text="备注")
        self.acc_tree.heading("phone", text="手机号")
        self.acc_tree.column("label", width=240, anchor=tk.W)
        self.acc_tree.column("phone", width=300, anchor=tk.W)
        self.acc_tree.pack(fill=tk.BOTH, expand=True)

        # 加大行高，列表更易读
        style = ttk.Style()
        try:
            style.configure("Acc.Treeview", rowheight=32, font=("Microsoft YaHei UI", 12))
            style.configure("Acc.Treeview.Heading", font=("Microsoft YaHei UI", 12, "bold"))
            self.acc_tree.configure(style="Acc.Treeview")
        except tk.TclError:
            pass

        btn_frm = ttk.Frame(acc_frm)
        btn_frm.pack(fill=tk.X, pady=(10, 0))
        self.launch_sel_btn = ttk.Button(
            btn_frm,
            text="一键启动选中",
            command=self.launch_selected_account,
            width=16,
        )
        self.launch_sel_btn.pack(side=tk.LEFT)
        self.launch_all_btn = ttk.Button(
            btn_frm,
            text="一键启动所有",
            command=self.launch_all_accounts,
            width=16,
        )
        self.launch_all_btn.pack(side=tk.LEFT, padx=(8, 0))
        ttk.Button(btn_frm, text="批量录入", command=self.import_accounts_excel, width=12).pack(side=tk.RIGHT)
        ttk.Button(btn_frm, text="批量导出", command=self.export_accounts_excel, width=12).pack(
            side=tk.RIGHT, padx=(0, 8)
        )
        ttk.Button(btn_frm, text="录入账号", command=self.add_account, width=12).pack(side=tk.RIGHT, padx=(0, 8))
        ttk.Button(btn_frm, text="修改账号", command=self.edit_account, width=12).pack(side=tk.RIGHT, padx=(0, 8))
        ttk.Button(btn_frm, text="删除账号", command=self.remove_account, width=12).pack(side=tk.RIGHT, padx=(0, 8))

        tile_frm = ttk.Frame(acc_frm)
        tile_frm.pack(fill=tk.X, pady=(8, 0))
        self.auto_tile_var = tk.BooleanVar(
            value=bool(load_settings().get("auto_tile_windows", False))
        )
        ttk.Checkbutton(
            tile_frm,
            text="自动分窗（批量启动后阶梯摆窗：第1个左上角，之后每次 +80,+80；第11个起右移1200px；最多一次20个）",
            variable=self.auto_tile_var,
            command=self._persist_auto_tile,
        ).pack(anchor=tk.W)
        self.auto_open_helper_var = tk.BooleanVar(
            value=bool(load_settings().get("auto_open_helper", False))
        )
        ttk.Checkbutton(
            tile_frm,
            text="序章助手自动打开（一进游戏就开，再拉多控/召唤；约3秒后缩到右上角；已手动打开则跳过）",
            variable=self.auto_open_helper_var,
            command=self._persist_auto_open_helper,
        ).pack(anchor=tk.W, pady=(4, 0))

        status_bar = ttk.Frame(outer)
        status_bar.pack(fill=tk.X, pady=(8, 0))
        self.status_var = tk.StringVar(value="就绪")
        ttk.Label(status_bar, textvariable=self.status_var, foreground="#666").pack(side=tk.LEFT)

        self.reload_accounts()

    def _set_status(self, text: str) -> None:
        self.root.after(0, lambda: self.status_var.set(text))

    def _persist_auto_tile(self) -> None:
        cfg = load_settings()
        cfg["auto_tile_windows"] = bool(self.auto_tile_var.get())
        try:
            save_settings(cfg)
        except OSError:
            pass

    def _persist_auto_open_helper(self) -> None:
        cfg = load_settings()
        cfg["auto_open_helper"] = bool(self.auto_open_helper_var.get())
        try:
            save_settings(cfg)
        except OSError:
            pass

    def _want_auto_open_helper(self) -> bool:
        try:
            return bool(self.auto_open_helper_var.get())
        except tk.TclError:
            return bool(load_settings().get("auto_open_helper", False))

    def _maybe_open_helper(self, instance_id: str, label: str) -> None:
        if not self._want_auto_open_helper():
            return
        try:
            ipc.open_helper_minimized(instance_id)
            self._set_status(f"[{label}] 已请求打开序章助手（约3秒后缩到右上角）")
        except Exception as exc:
            self._set_status(f"[{label}] 打开序章助手失败: {exc}")

    def pick_game_dir(self) -> None:
        chosen = filedialog.askdirectory(title="选择游戏根目录（含 cg37_Data）")
        if not chosen:
            return
        path = Path(chosen)
        if not (path / "cg37_Data").is_dir():
            messagebox.showerror("无效", "所选目录下没有 cg37_Data")
            return
        set_game_root(path)
        self.game_root_var.set(str(path))
        self._set_status(f"游戏目录已设为: {path}")

    def _game_root(self) -> Path:
        return Path(self.game_root_var.get().strip())

    def _bridge_is_ready(self) -> bool:
        root = self._game_root()
        if not root.is_dir():
            return False
        return is_mini_bridge_ready(root)

    def _warn_if_bridge_missing(self) -> bool:
        if self._bridge_is_ready():
            return True
        return messagebox.askyesno(
            "精简桥接未注入",
            "当前未检测到精简桥接，多开器启动后无法自动登录/召唤。\n\n"
            "请先在「序章补丁」中勾选「注入精简桥接」并应用。\n\n仍要启动游戏吗？",
        )

    def reload_accounts(self) -> None:
        for item in self.acc_tree.get_children():
            self.acc_tree.delete(item)
        for acc in load_accounts():
            self.acc_tree.insert("", tk.END, iid=acc.id, values=(acc.label, acc.phone))
        self._set_status(f"账号库共 {len(load_accounts())} 个账号")

    def add_account(self) -> None:
        label = simpledialog.askstring("备注", "账号备注（可选）:", parent=self.root) or ""
        phone = simpledialog.askstring("手机号", "手机号:", parent=self.root)
        if not phone:
            return
        password = simpledialog.askstring("密码", "密码:", show="*", parent=self.root)
        if password is None:
            return
        upsert_account(AccountProfile.create(label, phone, password))
        self.reload_accounts()

    def export_accounts_excel(self) -> None:
        """批量导出当前账号。表格式与批量录入相同，导出的文件可直接再录入。"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            messagebox.showerror(
                "缺少依赖",
                "需要 openpyxl 才能批量导出。\n请先执行：pip install openpyxl",
                parent=self.root,
            )
            return

        path = filedialog.asksaveasfilename(
            title="批量导出账号",
            defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")],
            initialfile="多开器账号批量导出.xlsx",
            parent=self.root,
        )
        if not path:
            return

        seen: set[str] = set()
        rows: list[tuple[str, str, str, str]] = []
        for acc in load_accounts():
            phone = (acc.phone or "").strip()
            if not phone or phone in seen:
                continue
            seen.add(phone)
            rows.append(
                (
                    (acc.label or phone).strip(),
                    phone,
                    acc.password or "",
                    acc.secondary_code or "",
                )
            )

        wb = Workbook()
        ws = wb.active
        ws.title = "账号"
        headers = ("备注", "手机号", "密码", "二级码")
        ws.append(list(headers))
        header_font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1F4E79")
        for col, _ in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for label, phone, password, sec in rows:
            ws.append([label, phone, password, sec])
            row_i = ws.max_row
            for col in (1, 2, 3, 4):
                cell = ws.cell(row=row_i, column=col)
                cell.number_format = "@"
                cell.alignment = Alignment(vertical="center")
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 12

        tip = wb.create_sheet("说明", 0)
        tip.column_dimensions["A"].width = 78
        tip["A1"] = "多开器账号批量导出"
        tip["A1"].font = Font(name="微软雅黑", bold=True, size=14)
        tips = [
            "本文件由「批量导出」生成，可直接用多开器或中控的「批量录入」导入。",
            "1. 「账号」表列：备注、手机号、密码、二级码。表头不要改名。",
            "2. 二级码可空；填写则必须是 6 位数字。",
            "3. 手机号相同视为同一账号：已有则更新，没有则新增。",
            "4. 备注/密码/二级码与库里不一致时，以本表录入的内容为准。",
            "5. 同一文件里手机号重复时，靠后的那一行生效。",
            "6. 空行、手机号或密码为空的行会跳过。",
        ]
        for i, line in enumerate(tips, start=3):
            tip[f"A{i}"] = line

        try:
            wb.save(path)
        except OSError as exc:
            messagebox.showerror("保存失败", str(exc), parent=self.root)
            return

        self._set_status(f"已批量导出 {len(rows)} 个账号：{path}")
        messagebox.showinfo(
            "完成",
            f"已导出 {len(rows)} 个账号：\n{path}\n\n该文件可直接用于「批量录入」。",
            parent=self.root,
        )

    def import_accounts_excel(self) -> None:
        """从 Excel 批量录入账号（同手机号则更新）。"""
        try:
            from openpyxl import load_workbook
        except ImportError:
            messagebox.showerror(
                "缺少依赖",
                "需要 openpyxl 才能批量录入。\n请先执行：pip install openpyxl",
                parent=self.root,
            )
            return

        path = filedialog.askopenfilename(
            title="选择账号 Excel",
            filetypes=[("Excel 工作簿", "*.xlsx"), ("所有文件", "*.*")],
            parent=self.root,
        )
        if not path:
            return

        try:
            added, updated, skipped, errors = self._import_accounts_from_xlsx(path)
        except Exception as exc:
            messagebox.showerror("录入失败", str(exc), parent=self.root)
            return

        self.reload_accounts()
        msg = f"新增 {added}，更新 {updated}，跳过 {skipped}"
        if errors:
            detail = "\n".join(errors[:12])
            if len(errors) > 12:
                detail += f"\n…另有 {len(errors) - 12} 条"
            messagebox.showwarning("批量录入完成（有跳过）", f"{msg}\n\n{detail}", parent=self.root)
        else:
            messagebox.showinfo("批量录入完成", msg, parent=self.root)
        self._set_status(f"批量录入：{msg}")

    @staticmethod
    def _import_accounts_from_xlsx(path: str) -> tuple[int, int, int, list[str]]:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = None
        for name in ("账号", "accounts", "Accounts"):
            if name in wb.sheetnames:
                ws = wb[name]
                break
        if ws is None:
            ws = wb[wb.sheetnames[0]]

        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            wb.close()
            raise ValueError("Excel 为空")

        def norm(cell: object) -> str:
            if cell is None:
                return ""
            if isinstance(cell, float) and cell.is_integer():
                return str(int(cell))
            if isinstance(cell, int):
                return str(cell)
            return str(cell).strip()

        col_map: dict[str, int] = {}
        aliases = {
            "label": ("备注", "账号备注", "名称", "label", "name"),
            "phone": ("手机号", "手机", "账号", "帐号", "phone", "user", "username"),
            "password": ("密码", "password", "pwd", "pass"),
            "secondary_code": ("二级码", "二级密码", "安全锁", "secondary_code", "sec", "code"),
        }
        for idx, raw in enumerate(header or ()):
            key = norm(raw).lower()
            for field, names in aliases.items():
                if key in {n.lower() for n in names} and field not in col_map:
                    col_map[field] = idx
        if "phone" not in col_map or "password" not in col_map:
            wb.close()
            raise ValueError("表头需包含「手机号」「密码」列（备注/二级码可选）")

        existing_by_phone: dict[str, AccountProfile] = {}
        for acc in load_accounts():
            phone_key = (acc.phone or "").strip()
            if phone_key and phone_key not in existing_by_phone:
                existing_by_phone[phone_key] = acc
        pending: dict[str, tuple[str, str, str, bool]] = {}
        added = updated = skipped = 0
        errors: list[str] = []
        for row_i, row in enumerate(rows, start=2):
            if row is None:
                continue
            cells = list(row)
            phone = norm(cells[col_map["phone"]]) if col_map["phone"] < len(cells) else ""
            password = norm(cells[col_map["password"]]) if col_map["password"] < len(cells) else ""
            label = ""
            if "label" in col_map and col_map["label"] < len(cells):
                label = norm(cells[col_map["label"]])
            sec_raw = ""
            has_sec_col = "secondary_code" in col_map
            if has_sec_col and col_map["secondary_code"] < len(cells):
                sec_raw = norm(cells[col_map["secondary_code"]])
            if not phone and not password and not label and not sec_raw:
                continue
            if not phone or not password:
                skipped += 1
                errors.append(f"第{row_i}行：手机号或密码为空，已跳过")
                continue
            sec = ""
            if has_sec_col:
                sec, err = normalize_secondary_code(sec_raw)
                if err:
                    skipped += 1
                    errors.append(f"第{row_i}行：{err}，已跳过")
                    continue
            pending[phone] = (label, password, sec, has_sec_col)

        for phone, (label, password, sec, has_sec_col) in pending.items():
            label_use = label or phone
            if phone in existing_by_phone:
                acc = existing_by_phone[phone]
                same = (
                    acc.label == label_use
                    and acc.password == password
                    and (not has_sec_col or acc.secondary_code == sec)
                )
                if same:
                    continue
                acc.label = label_use
                acc.password = password
                if has_sec_col:
                    acc.secondary_code = sec
                upsert_account(acc)
                updated += 1
            else:
                acc = AccountProfile.create(
                    label, phone, password, secondary_code=sec if has_sec_col else ""
                )
                upsert_account(acc)
                existing_by_phone[phone] = acc
                added += 1

        wb.close()
        return added, updated, skipped, errors

    def edit_account(self) -> None:
        sel = self.acc_tree.selection()
        if not sel:
            messagebox.showwarning("未选择", "请先选择要修改的账号")
            return
        acc = next((a for a in load_accounts() if a.id == sel[0]), None)
        if acc is None:
            return
        label = simpledialog.askstring("备注", "账号备注（可选）:", initialvalue=acc.label, parent=self.root)
        if label is None:
            return
        phone = simpledialog.askstring("手机号", "手机号:", initialvalue=acc.phone, parent=self.root)
        if not phone:
            return
        password = simpledialog.askstring("密码", "密码:", initialvalue=acc.password, show="*", parent=self.root)
        if password is None:
            return
        acc.label = label.strip() or phone
        acc.phone = phone.strip()
        acc.password = password
        upsert_account(acc)
        self.reload_accounts()

    def remove_account(self) -> None:
        sel = self.acc_tree.selection()
        if not sel:
            messagebox.showwarning("未选择", "请先选择要删除的账号")
            return
        acc = next((a for a in load_accounts() if a.id == sel[0]), None)
        name = acc.label if acc else sel[0]
        if messagebox.askyesno("确认", f"删除账号「{name}」？"):
            delete_account(sel[0])
            self.reload_accounts()

    def launch_selected_account(self) -> None:
        sel = self.acc_tree.selection()
        if not sel:
            messagebox.showwarning("未选择", "请先选择账号")
            return
        if not self._warn_if_bridge_missing():
            return
        accounts = load_accounts()
        by_id = {a.id: a for a in accounts}
        chosen = [by_id[i] for i in sel if i in by_id]
        if not chosen:
            return
        self._launch_staggered(chosen, "启动选中")

    def launch_all_accounts(self) -> None:
        accounts = load_accounts()
        if not accounts:
            messagebox.showwarning("无账号", "账号库为空，请先录入账号。")
            return
        if not self._warn_if_bridge_missing():
            return
        self._launch_staggered(accounts, "一键启动所有")

    # --- 逐个启动（间隔 10 秒，避免同时初始化黑屏） ---

    LAUNCH_GAP_SEC = 10

    def _launch_staggered(self, accounts: list[AccountProfile], label: str) -> None:
        """按账号库顺序逐个启动，每个间隔 LAUNCH_GAP_SEC 秒，避免多开同时初始化黑屏。"""
        tile = bool(self.auto_tile_var.get())
        chosen = list(accounts)
        if tile and len(chosen) > TILE_MAX_WINDOWS:
            extra = len(chosen) - TILE_MAX_WINDOWS
            if not messagebox.askyesno(
                "自动分窗",
                f"自动分窗一次最多开 {TILE_MAX_WINDOWS} 个窗口。\n"
                f"当前选了 {len(chosen)} 个，将只启动前 {TILE_MAX_WINDOWS} 个（跳过 {extra} 个）。\n\n继续？",
            ):
                return
            chosen = chosen[:TILE_MAX_WINDOWS]
        threading.Thread(
            target=self._launch_staggered_worker,
            args=(chosen, label, tile),
            daemon=True,
        ).start()

    def _place_tiled_window(self, pid: int, slot: int, name: str) -> None:
        x, y = _tile_xy(slot)
        hwnd = _wait_hwnd_for_pid(pid)
        if not hwnd:
            self._set_status(f"[{name}] 窗口未出现，跳过分窗 ({x},{y})")
            return
        moved = _move_window(hwnd, x, y)
        # Unity 启动画面关掉后会换主窗口，再摆几次以免弹回默认位置
        for _ in range(4):
            time.sleep(1.2)
            again = _find_hwnd_for_pid(pid)
            if again:
                hwnd = again
                moved = _move_window(hwnd, x, y) or moved
        if moved:
            self._set_status(f"[{name}] 已分窗到 ({x},{y})")
        else:
            self._set_status(f"[{name}] 分窗失败 ({x},{y})")

    def _launch_staggered_worker(
        self, accounts: list[AccountProfile], label: str, tile: bool
    ) -> None:
        ok = 0
        errors: list[str] = []
        total = len(accounts)
        for idx, acc in enumerate(accounts, start=1):
            name = acc.label or acc.phone
            try:
                inst = launch_game(self._game_root())
                self.instances.append(inst)
                threading.Thread(
                    target=self._auto_workflow,
                    args=(inst.instance_id, acc.phone, acc.password, name),
                    daemon=True,
                ).start()
                ok += 1
                self._set_status(f"{label}进行中：{name}（{idx}/{total}）已启动")
                if tile:
                    self._place_tiled_window(inst.pid, idx - 1, name)
            except Exception as exc:
                errors.append(f"{name}: {exc}")

            if idx < total:
                self._set_status(f"{label}进行中：{idx}/{total} 已启动，{self.LAUNCH_GAP_SEC}秒后启动下一个…")
                time.sleep(self.LAUNCH_GAP_SEC)

        tile_part = "，已自动分窗" if tile else ""
        summary = f"{label}完成：成功 {ok}/{total}（已自动进入登录→拉多控→召唤流程，逐个间隔{self.LAUNCH_GAP_SEC}秒{tile_part}）"
        if errors:
            summary += "\n" + "\n".join(errors[:10])
        self._set_status(summary)
        if errors or label == "一键启动所有":
            self.root.after(0, lambda: messagebox.showwarning(label, summary, parent=self.root))

    def _auto_workflow(self, instance_id: str, phone: str, password: str, label: str) -> None:
        """启动后后台自动走精简桥接 workflow_step1：等桥接→登录→进游戏→拉多控→一键召唤。"""
        try:
            self._set_status(f"[{label}] 等待精简桥接连接…")
            if not ipc.wait_for_bridge(instance_id, timeout=240):
                self._set_status(f"[{label}] 桥接连接超时（未注入精简桥接？）")
                return
            self._set_status(f"[{label}] 桥接已连接，自动登录/开助手/拉多控/召唤…")
            ipc.workflow_step1_five_chars(
                instance_id,
                phone,
                password,
                open_helper=self._want_auto_open_helper(),
            )
            ok, msg = ipc.wait_workflow_done(instance_id, timeout=600)
            if ok:
                self._set_status(f"[{label}] 流程完成")
            else:
                self._set_status(f"[{label}] 流程失败: {msg}")
        except Exception as exc:
            self._set_status(f"[{label}] 异常: {type(exc).__name__}: {exc}")

    # --- 批量一键（协议驱动，靠 bridge 命令与 team/multi_ready 判定） ---

    def _live_bridge_instances(self) -> list[str]:
        """返回当前有桥接心跳的实例 ID（按 PID 稳定排序）。"""
        rows = ipc.list_instance_snapshots()
        live = [r for r in rows if r.get("alive")]
        live.sort(key=lambda r: r.get("pid_txt") or 0)
        return [r["instance_id"] for r in live]

    def _require_accounts(self) -> list[AccountProfile] | None:
        accounts = load_accounts()
        if not accounts:
            messagebox.showwarning("无账号", "账号库为空，请先录入账号。")
            return None
        return accounts

    def batch_login_fetch(self) -> None:
        """批量：按账号库顺序对每个已注入桥接的实例 登录→进游戏→拉起离线多控。"""
        if not self._warn_if_bridge_missing():
            return
        accounts = self._require_accounts()
        if accounts is None:
            return
        iids = self._live_bridge_instances()
        if not iids:
            messagebox.showwarning(
                "无实例",
                "没有检测到已注入精简桥接的实例。\n请先启动游戏并确保已注入精简桥接。",
            )
            return
        if len(iids) > len(accounts):
            messagebox.showwarning(
                "账号不足",
                f"当前 {len(iids)} 个实例，但账号库只有 {len(accounts)} 个账号。\n"
                "多余实例将跳过登录。",
            )
        pairs = list(zip(iids, accounts))
        self._set_batch_busy(True, f"正在批量登录拉取 {len(pairs)} 个实例…")
        threading.Thread(target=self._batch_login_worker, args=(pairs,), daemon=True).start()

    def _batch_login_worker(self, pairs: list[tuple[str, AccountProfile]]) -> None:
        ok_count = 0
        lines: list[str] = []
        for iid, acc in pairs:
            label = acc.label or acc.phone
            st = ipc.read_state(iid) or {}
            phase = st.get("phase", "")
            try:
                if int(st.get("team_num") or 0) >= 2:
                    ok_count += 1
                    lines.append(f"[OK] {label}: 已在队伍中，跳过拉取")
                    self._set_batch_status(f"[{label}] 已组队，视为成功")
                    self._maybe_open_helper(iid, label)
                    continue
                if phase == "in_game":
                    self._set_batch_status(f"[{label}] 已进游戏，先开助手再拉多控…")
                    self._maybe_open_helper(iid, label)
                    ipc.multi_login_offline_all(iid)
                    multi_ok = ipc.wait_multi_ready(iid, timeout=120)
                else:
                    self._set_batch_status(f"[{label}] 登录→进游戏(开助手)→拉起多控…")
                    ipc.workflow_login_enter(
                        iid,
                        acc.phone,
                        acc.password,
                        open_helper=self._want_auto_open_helper(),
                    )
                    entered = ipc.wait_for_in_game(iid, timeout=300)
                    if not entered:
                        ok, msg = ipc.wait_workflow_done(iid, timeout=60)
                        if not ok:
                            lines.append(f"[FAIL] {label}: 登录/进游戏失败 {msg}")
                            self._set_batch_status(f"[{label}] 失败")
                            continue
                    after = ipc.read_state(iid) or {}
                    if int(after.get("team_num") or 0) >= 2:
                        ok_count += 1
                        lines.append(f"[OK] {label}: 已在队伍中，跳过拉取")
                        self._set_batch_status(f"[{label}] 已组队，视为成功")
                        continue
                    ipc.multi_login_offline_all(iid)
                    multi_ok = ipc.wait_multi_ready(iid, timeout=120)
                if multi_ok:
                    ok_count += 1
                    team_n = int((ipc.read_state(iid) or {}).get("team_num") or 0)
                    if team_n >= 2:
                        lines.append(f"[OK] {label}: 已在队伍中")
                    else:
                        lines.append(f"[OK] {label}: 已进游戏并拉起多控")
                    self._set_batch_status(f"[{label}] 完成")
                else:
                    lines.append(f"[WARN] {label}: 已进游戏但多控未全部上线（可稍后一键召唤）")
                    self._set_batch_status(f"[{label}] 多控未全上线")
            except Exception as exc:
                lines.append(f"[FAIL] {label}: {type(exc).__name__}: {exc}")
                self._set_batch_status(f"[{label}] 异常")
        summary = f"批量登录拉取完成：成功 {ok_count}/{len(pairs)}"
        if lines:
            summary += "\n" + "\n".join(lines)
        self._set_batch_done(summary)

    def batch_summon(self) -> None:
        """批量：对所有实例发一键召唤（协议），按 team≥5 判定聚齐。"""
        iids = self._live_bridge_instances()
        if not iids:
            messagebox.showwarning("无实例", "没有检测到已注入桥接的实例。")
            return
        self._set_batch_busy(True, f"正在一键召唤 {len(iids)} 个实例…")
        threading.Thread(target=self._batch_summon_worker, args=(iids,), daemon=True).start()

    def _batch_summon_worker(self, iids: list[str]) -> None:
        ok_count = 0
        lines: list[str] = []
        for iid in iids:
            st = ipc.read_state(iid) or {}
            if int(st.get("team_num") or 0) >= 2:
                ok_count += 1
                lines.append(f"[OK] {iid}: 已在队伍中，跳过召唤")
                self._maybe_open_helper(iid, iid)
                continue
            self._set_batch_status(f"[{iid}] 一键召唤…")
            self._maybe_open_helper(iid, iid)
            ipc.one_key_summon(iid)
            ok = ipc.wait_for_team(iid, timeout=90)
            if ok:
                ok_count += 1
                lines.append(f"[OK] {iid}: 队伍聚齐 (team≥5)")
            else:
                # 补一发队伍召集再试
                ipc.team_gather(iid)
                ok2 = ipc.wait_for_team(iid, timeout=60)
                if ok2:
                    ok_count += 1
                    lines.append(f"[OK] {iid}: 队伍召集后聚齐")
                else:
                    lines.append(f"[FAIL] {iid}: 召唤/召集未聚齐")
        summary = f"一键召唤完成：成功 {ok_count}/{len(iids)}"
        if lines:
            summary += "\n" + "\n".join(lines)
        self._set_batch_done(summary)

    def launch_window_monitor(self) -> None:
        """启动窗口监视（与傻瓜补丁「启动窗口监视」同入口）。"""
        cands = self._window_monitor_candidates()
        if not cands:
            messagebox.showinfo(
                APP_TITLE,
                "未找到窗口监视脚本。\n\n"
                "预期路径：魔力宝贝序章补丁\\scripts\\window_monitor_gui.py\n"
                "或包内「窗口监视.exe」。",
                parent=self.root,
            )
            return
        try:
            exe = cands[0]
            if exe.suffix.lower() == ".py":
                pyw = Path(sys.executable).with_name("pythonw.exe")
                py = str(pyw) if pyw.is_file() else sys.executable
                subprocess.Popen([py, str(exe)], cwd=str(exe.parent))
            else:
                subprocess.Popen([str(exe)], cwd=str(exe.parent))
            self.status_var.set(f"已启动窗口监视：{exe.name}")
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f"无法启动窗口监视：\n{exc}", parent=self.root)

    def _window_monitor_candidates(self) -> list[Path]:
        cands: list[Path] = []
        here = Path(__file__).resolve()
        # 开发：仓库内补丁 scripts；发布：与多开器同级的窗口监视.exe
        cands.extend(
            [
                here.parents[2] / "魔力宝贝序章补丁" / "scripts" / "window_monitor_gui.py",
                here.parent / "window_monitor_gui.py",
                here.parents[1] / "窗口监视.exe",
                here.parents[1] / "window_monitor_gui.exe",
            ]
        )
        if getattr(sys, "frozen", False):
            exe_dir = Path(sys.executable).resolve().parent
            cands.extend(
                [
                    exe_dir / "窗口监视.exe",
                    exe_dir / "窗口监视" / "窗口监视.exe",
                    exe_dir / "window_monitor_gui.exe",
                ]
            )
        return [p for p in cands if p.is_file()]

    def _set_batch_busy(self, busy: bool, status: str) -> None:
        state = ("disabled" if busy else "normal")
        for btn in (self.batch_login_btn, self.batch_summon_btn):
            btn.config(state=state)
        self.batch_status_var.set(status)

    def _set_batch_status(self, text: str) -> None:
        self.root.after(0, lambda: self.batch_status_var.set(text))

    def _set_batch_done(self, text: str) -> None:
        def _finish() -> None:
            self._set_batch_busy(False, "批量状态：就绪")
            self.batch_status_var.set(text)
        self.root.after(0, _finish)


def find_account_by_query(query: str) -> AccountProfile:
    q = (query or "").strip()
    if not q:
        raise ValueError("账号名为空")
    accounts = load_accounts()
    hits = [a for a in accounts if a.label == q or a.phone == q]
    if not hits:
        hits = [a for a in accounts if q in (a.label or "")]
    if not hits:
        names = "、".join((a.label or a.phone) for a in accounts) or "(空)"
        raise ValueError(f"账号库没有「{q}」。现有：{names}")
    if len(hits) > 1:
        names = "、".join(a.label or a.phone for a in hits)
        raise ValueError(f"「{q}」匹配到多个账号：{names}")
    return hits[0]


def cli_launch_and_login(label: str, *, login_only: bool = False) -> int:
    """不打开 GUI：启动一个号并一键登录。默认登录后拉离线多控并一键召唤。"""
    acc = find_account_by_query(label)
    root = get_game_root()
    name = acc.label or acc.phone
    phone_mask = (acc.phone[:3] + "***") if acc.phone else ""
    print(f"[INFO] 游戏目录 {root}", flush=True)
    print(f"[INFO] 账号 {name} {phone_mask}", flush=True)
    if not is_mini_bridge_ready(root):
        print("[FAIL] 当前游戏目录未注入精简桥接，无法自动登录", flush=True)
        return 1

    inst = launch_game(root)
    print(f"[OK] 已启动 pid={inst.pid} instance_id={inst.instance_id}", flush=True)
    print(f"PID={inst.pid}", flush=True)

    print(f"[INFO] [{name}] 等待精简桥接…", flush=True)
    if not ipc.wait_for_bridge(inst.instance_id, timeout=240):
        print(f"[FAIL] [{name}] 桥接连接超时 pid={inst.pid}", flush=True)
        return 2

    if login_only:
        print(f"[INFO] [{name}] 一键登录…", flush=True)
        want_helper = bool(load_settings().get("auto_open_helper"))
        ipc.workflow_login_enter(
            inst.instance_id,
            acc.phone,
            acc.password,
            open_helper=want_helper,
        )
        ok, msg, _st = ipc.wait_for_in_game(inst.instance_id, timeout=300)
        if ok:
            uid = str(msg or "")
            tail = uid[-4:] if len(uid) > 4 else uid
            print(f"[OK] [{name}] 已进游戏 pid={inst.pid} uid尾={tail}", flush=True)
            return 0
        print(f"[FAIL] [{name}] 登录失败 {msg} pid={inst.pid}", flush=True)
        return 3

    print(f"[INFO] [{name}] 登录→开助手→拉多控→召唤…", flush=True)
    want_helper = bool(load_settings().get("auto_open_helper"))
    ipc.workflow_step1_five_chars(
        inst.instance_id,
        acc.phone,
        acc.password,
        open_helper=want_helper,
    )
    ok, msg = ipc.wait_workflow_done(inst.instance_id, timeout=600)
    if ok:
        print(f"[OK] [{name}] 流程完成 pid={inst.pid}", flush=True)
        return 0
    print(f"[FAIL] [{name}] {msg} pid={inst.pid}", flush=True)
    return 3


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=APP_TITLE)
    parser.add_argument(
        "--launch",
        metavar="账号备注",
        help="启动指定账号：一键登录并拉多控+召唤（不打开 GUI）",
    )
    parser.add_argument(
        "--login-only",
        action="store_true",
        help="只登录进游戏，不拉离线多控、不召唤",
    )
    args = parser.parse_args(argv)
    if args.launch:
        try:
            return cli_launch_and_login(args.launch, login_only=args.login_only)
        except Exception as exc:
            print(f"[FAIL] {type(exc).__name__}: {exc}", flush=True)
            return 1
    if not ensure_single_instance(
        APP_TITLE,
        message=(
            f"{APP_TITLE} 已在运行，不能重复打开。\n\n"
            "本程序仅允许运行一个窗口；请在已打开的窗口中管理游戏多开。"
        ),
    ):
        return 1
    app = MultiLauncherApp()
    app.root.mainloop()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
